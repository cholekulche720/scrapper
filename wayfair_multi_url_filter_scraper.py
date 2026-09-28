#!/usr/bin/env python3
"""
Wayfair Multi-URL Filter & Category Hierarchy Scraper
=====================================================
Recursively traverses a list of Wayfair category/browse URLs:
1. Traverses from each starting URL down through all intermediate subcategories
   to the deepest leaf listing nodes (/sb0/ or product listing pages).
2. At every level:
   - Discovers child categories and records a 'Category Navigation' row under the parent.
   - Extracts direct product filter blocks (facets, checkboxes, ranges, swatches, etc.)
     from the page if filters are present.
3. Recursively descends into all children down to the leaf node.
4. Prevents duplicate scraping across multiple URLs that share child categories.
5. Exports the complete hierarchy and filter dataset to a richly styled Excel (.xlsx) file
   with formatted column headers (e.g., 'Category Name', 'Category Url'),
   category separator borders, color-coded rows, and auto-adjusted column widths.
"""

import argparse
import base64
import csv
import json
import os
import re
import sys
import time
from typing import Any, Dict, List, Optional, Set, Tuple
import requests
from bs4 import BeautifulSoup

# ==============================================================================
# DEFAULT SEED URLS (You can paste your list of URLs here)
# ==============================================================================
DEFAULT_SEED_URLS: List[str] = [
    # Examples (replace or add as many URLs as you need):
    "https://www.wayfair.com/appliances/cat/appliances-c215602.html",
]

DEFAULT_ZYTE_KEY = "62eb8bf11a3e433aad5a8cdbcbccfe69"
DEFAULT_EXCEL_OUTPUT = "wayfair_appliances_filters.xlsx"
DEFAULT_CSV_OUTPUT = "wayfair_appliances_filters.csv"

# Internal field names (dictionary keys)
FIELDNAMES = [
    "category_name",
    "category_url",
    "hierarchy_path",
    "section",
    "filter_group",
    "sub_group",
    "attribute_name",
    "attribute_id",
    "filter_id",
    "input_type",
    "description",
    "tooltip",
    "range_min",
    "range_max",
    "range_unit",
    "is_enabled",
    "is_selected",
    "link_url",
]

# Proper human-readable headers with spacing for the generated Excel file
EXCEL_HEADER_MAPPING: Dict[str, str] = {
    "category_name": "Category Name",
    "category_url": "Category Url",
    "hierarchy_path": "Hierarchy Path",
    "section": "Section",
    "filter_group": "Attribute",
    "sub_group": "Sub Attribute",
    "attribute_name": "Options",
    "attribute_id": "Attribute Id",
    "filter_id": "Filter Id",
    "input_type": "Input Type",
    "description": "Description",
    "tooltip": "Tooltip",
    "range_min": "Range Min",
    "range_max": "Range Max",
    "range_unit": "Range Unit",
    "is_enabled": "Is Enabled",
    "is_selected": "Is Selected",
    "link_url": "Link Url",
}


def is_sale_category(name: str, url: str, ignore_sales: bool = False) -> bool:
    """
    Checks whether a category name or URL is related to sales or clearance.
    Returns True to skip if ignore_sales is enabled.
    """
    if not ignore_sales:
        return False

    name_lower = (name or "").lower().strip()
    url_lower = (url or "").lower().strip()

    sale_keywords = [
        "sale", "sales", "clearance", "warehouse clearout", "clearout",
        "daily sales", "daily deals", "deals", "closeout", "closeouts",
        "outlet", "open box", "liquidation", "special buy", "discount",
    ]

    for kw in sale_keywords:
        if re.search(r"\b" + re.escape(kw) + r"\b", name_lower):
            return True
        if kw in url_lower:
            return True

    return False


TOP_BAR_CATEGORIES = [
    ("Furniture", "https://www.wayfair.com/furniture/cat/furniture-c45974.html"),
    ("Outdoor", "https://www.wayfair.com/outdoor/cat/outdoor-c32334.html"),
    ("Bedding & Bath", "https://www.wayfair.com/bed-bath/cat/bed-bath-c215329.html"),
    ("Rugs", "https://www.wayfair.com/rugs/cat/rugs-c215385.html"),
    ("Decor & Pillows", "https://www.wayfair.com/decor-pillows/cat/decor-pillows-c45752.html"),
    ("Lighting", "https://www.wayfair.com/lighting/cat/lighting-c215735.html"),
    ("Organization", "https://www.wayfair.com/storage-organization/cat/storage-organization-c215875.html"),
    ("Kitchen", "https://www.wayfair.com/kitchen-tabletop/cat/kitchen-tabletop-c45667.html"),
    ("Baby & Kids", "https://www.wayfair.com/baby-kids/cat/baby-kids-c45226.html"),
    ("Home Improvement", "https://www.wayfair.com/home-improvement/cat/home-improvement-c524317.html"),
    ("Appliances", "https://www.wayfair.com/appliances/cat/appliances-c215602.html"),
    ("Pet", "https://www.wayfair.com/pet/cat/pet-c533276.html"),
    ("Holiday", "https://www.wayfair.com/holiday-decor/cat/holiday-decor-c47391.html"),
    ("Shop by Room", "https://www.wayfair.com/shop-by-room/cat/shop-by-room-c1876502.html"),
]
TOP_BAR_URLS = {url.split("?")[0] for _, url in TOP_BAR_CATEGORIES}


# ==============================================================================
# COLOR PALETTES & ROW STYLING FOR EXCEL
# ==============================================================================
CATEGORY_NAV_BG = "E0F2FE"       # Light sky blue
FEATURED_FILTER_BG = "FEF3C7"    # Soft warm amber

PALETTE_ROULETTE = [
    ("EDE9FE", "4C1D95"),  # Soft Purple
    ("E0E7FF", "3730A3"),  # Soft Indigo
    ("ECFCCB", "365314"),  # Soft Lime
    ("FCE7F3", "831843"),  # Soft Pink
    ("CCFBF1", "115E59"),  # Soft Teal
    ("FFEDD5", "7C2D12"),  # Soft Peach
    ("CFFAFE", "155E75"),  # Soft Cyan
    ("F3E8FF", "581C87"),  # Soft Violet
]

MATERIAL_PALETTE = {
    "leather": ("F5EBE1", "5C3A21"),
    "fabric": ("E8F0E6", "2B4436"),
    "velvet": ("F3E8EE", "612940"),
    "wood": ("EDE4DC", "4A3326"),
    "metal": ("E5E7EB", "1F2937"),
}

COLOR_PALETTE = {
    "upholstery color": ("EBE4EB", "4A2840"),
    "frame color": ("E0EBEB", "234B4B"),
    "wood tone": ("EBE6DF", "4D3C2C"),
}


def get_row_color_scheme(
    row_data: Dict[str, str], group_color_cache: Dict[str, int]
) -> Tuple[str, str]:
    """Returns (background_hex, text_hex) styling based on section and group."""
    section = row_data.get("section", "")
    group = (row_data.get("filter_group", "") or "").strip().lower()
    sub_group = (row_data.get("sub_group", "") or "").strip().lower()

    if section == "Category Navigation":
        return CATEGORY_NAV_BG, "0369A1"

    if section == "Featured Filters":
        return FEATURED_FILTER_BG, "92400E"

    if group == "material":
        for mat_key, colors in MATERIAL_PALETTE.items():
            if mat_key in sub_group:
                return colors
        return "F5EBE1", "4A3326"

    if group == "color":
        for col_key, colors in COLOR_PALETTE.items():
            if col_key in sub_group:
                return colors
        return "EBE4EB", "3730A3"

    if group:
        if group not in group_color_cache:
            group_color_cache[group] = len(group_color_cache) % len(PALETTE_ROULETTE)
        idx = group_color_cache[group]
        return PALETTE_ROULETTE[idx]

    return "FFFFFF", "000000"


# ==============================================================================
# MAIN SCRAPER CLASS
# ==============================================================================
class WayfairMultiFilterScraper:
    """
    Crawls multiple Wayfair category pages down through all intermediate levels
    to leaf nodes, extracting category hierarchy navigation rows and all filter blocks.
    """

    def __init__(
        self,
        zyte_key: Optional[str] = None,
        use_zyte: bool = True,
        include_nav: bool = True,
        ignore_sales: bool = False,
        timeout: int = 45,
        delay: float = 1.0,
    ):
        self.zyte_key = zyte_key or os.getenv("ZYTE_API_KEY") or DEFAULT_ZYTE_KEY
        self.use_zyte = use_zyte and bool(self.zyte_key)
        self.include_nav = include_nav
        self.ignore_sales = ignore_sales
        self.timeout = timeout
        self.delay = delay
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
            ),
            "Accept-Language": "en-US,en;q=0.9",
        })

    def fetch_page_html(self, url: str) -> str:
        """Fetches page HTML with Zyte API anti-blocking proxy, with direct request fallback."""
        if self.use_zyte:
            print(f"[*] Fetching {url} via Zyte API...")
            try:
                api_url = "https://api.zyte.com/v1/extract"
                auth = (self.zyte_key, "")
                payload = {
                    "url": url,
                    "httpResponseBody": True,
                    "httpResponseHeaders": True,
                }
                response = requests.post(
                    api_url, auth=auth, json=payload, timeout=self.timeout
                )
                if response.status_code == 200:
                    data = response.json()
                    b64_body = data.get("httpResponseBody")
                    if b64_body:
                        html = base64.b64decode(b64_body).decode("utf-8", errors="ignore")
                        print(f"[+] Successfully fetched {len(html)} bytes via Zyte API")
                        return html
                print(f"[-] Zyte API returned status {response.status_code}: {response.text[:200]}")
                print("[!] Falling back to direct request...")
            except Exception as e:
                print(f"[-] Zyte API request failed: {e}. Falling back to direct request...")

        print(f"[*] Fetching {url} via direct HTTP request...")
        for attempt in range(2):
            try:
                response = self.session.get(url, timeout=self.timeout)
                if response.status_code == 200:
                    print(f"[+] Successfully fetched {len(response.text)} bytes via direct request")
                    return response.text
                print(f"[-] Direct request status: {response.status_code}")
            except Exception as ex:
                print(f"[-] Direct request attempt {attempt + 1} failed: {ex}")
            time.sleep(1)

        raise RuntimeError(f"Failed to fetch {url} via Zyte API and direct HTTP.")

    def extract_flight_stream(self, html: str) -> str:
        """Extracts and concatenates Next.js React Server Component flight stream."""
        scripts = re.findall(
            r"<script>(self\.__next_f\.push\(.*?\))</script>", html, re.DOTALL
        )
        stream_chunks = []
        for s in scripts:
            match = re.search(
                r'self\.__next_f\.push\(\[(\d+),\s*"((?:[^"\\]|\\.)*)"\]\)',
                s,
                re.DOTALL,
            )
            if match:
                c_type = int(match.group(1))
                c_raw = match.group(2)
                try:
                    c_text = json.loads(f'"{c_raw}"')
                except Exception:
                    c_text = c_raw.encode("utf-8").decode("unicode_escape")
                if c_type == 1:
                    stream_chunks.append(c_text)

        return "".join(stream_chunks)

    def discover_subcategories(self, parent_url: str, html: str) -> List[Dict[str, str]]:
        """
        Extracts subcategories from a category landing page.
        - If the page is a search-browse page (/sb0/), it is a leaf listing with product filters and no child categories.
        - On landing pages (/cat/), strips header, nav, footer, script, and style tags to prevent capturing global menus.
        - Detects legitimate category cards (CompactImageCard, DepartmentCard, CategoryCard, BrowseCard, Tile, etc.)
          and flight stream productCategory navigation blocks.
        - Filters out self-links, top-bar department links, and non-category paths (/shop-by-room/, /pdp/, etc.).
        """
        clean_parent = parent_url.split("?")[0]
        if "/sb0/" in clean_parent:
            return []

        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["header", "nav", "footer", "script", "style"]):
            tag.decompose()

        subcategories: List[Dict[str, str]] = []
        seen: Set[str] = set()

        parent_cid_match = re.search(r"-c(\d+)\.html", clean_parent)
        parent_cid = parent_cid_match.group(1) if parent_cid_match else ""

        # Method 1: Category cards in DOM
        for a in soup.find_all("a", href=True):
            clio_id = a.get("data-clio-id", "")
            href = a["href"].split("?")[0]
            if href.startswith("/"):
                href = f"https://www.wayfair.com{href}"

            if href in seen or href == clean_parent:
                continue

            # Prevent self-linking or looping back to parent category ID
            if parent_cid and f"-c{parent_cid}.html" in href:
                continue

            # Prevent linking to other top-bar departments
            if href in TOP_BAR_URLS:
                continue

            # Exclude non-category sections (room galleries, brands, account, pdp, etc.)
            if any(skip in href for skip in ["/shop-by-room/", "/brand/", "/account/", "/help/", "/services/", "/v/", "/pdp/"]):
                continue

            # Must be a category or search browse page
            if not ("/cat/" in href or "/sb0/" in href or "/daily-sales" in href or "/curated/" in href):
                continue

            is_card = any(k in clio_id for k in ["CompactImageCard", "DepartmentCard", "CategoryCard", "BrowseCard", "Tile", "Card"])
            if not is_card:
                continue

            img = a.find("img")
            img_alt = (img.get("alt") or "").strip() if img else ""
            link_text = a.get_text(strip=True)
            display_name = img_alt if img_alt else link_text

            if not display_name:
                continue

            if (
                is_sale_category(display_name, href, self.ignore_sales)
                or is_sale_category(link_text, href, self.ignore_sales)
                or is_sale_category(img_alt, href, self.ignore_sales)
            ):
                continue

            seen.add(href)
            subcategories.append({"name": display_name, "url": href})

        # Method 2: Curated navigation blocks in flight stream (fallback if no cards found in DOM)
        if not subcategories:
            pattern = (
                r"\"productCategory\":\{\"__typename\":\"ProductCategory\",\"categoryId\":\d+,"
                r"\"displayName\":\"([^\"]+)\",\"url\":\"(https://www\.wayfair\.com/[^\"/]+/cat/[^\"]+|https://www\.wayfair\.com/[^\"/]+/sb0/[^\"]+)\""
            )
            for name, url in re.findall(pattern, html):
                clean_url = url.split("?")[0]
                if clean_url not in seen and clean_url != clean_parent and (not parent_cid or f"-c{parent_cid}.html" not in clean_url):
                    if clean_url in TOP_BAR_URLS:
                        continue
                    if is_sale_category(name, clean_url, self.ignore_sales):
                        continue
                    seen.add(clean_url)
                    subcategories.append({"name": name, "url": clean_url})

        return subcategories

    def parse_flight_blocks(
        self, stream: str
    ) -> Tuple[Optional[Dict[str, Any]], Optional[List[Dict[str, Any]]], List[List[Dict[str, Any]]]]:
        """
        Locates and parses category navigation, featured filters, and main filter blocks.
        Supports:
        - BlockBuilderBrowseCategoryNav (breadcrumbs & siblings)
        - BlockBuilderBrowseFeaturedFilters (featured pills & toggles)
        - BlockBuilderBrowseFilterRefinements (legacy appliedFilters format)
        - BlockBuilderBrowseFilters (modern browse GraphQL format with FilterSubsections, FilterWithSelectableOptions, etc.)
        - BlockBuilderCategoryStandardExperience (category landing page GraphQL format with data.refinements.filters)
        """
        lines = stream.split("\n")
        nav_data = None
        featured_data: List[Dict[str, Any]] = []
        raw_filter_blocks: List[List[Dict[str, Any]]] = []

        for line in lines:
            if not line:
                continue
            colon_pos = line.find(":")
            if colon_pos == -1:
                continue

            payload_str = line[colon_pos + 1 :]

            # 1. Category Navigation
            if "BlockBuilderBrowseCategoryNav" in line and not nav_data:
                try:
                    parsed_nav = json.loads(payload_str)
                    if isinstance(parsed_nav, list) and len(parsed_nav) > 3 and isinstance(parsed_nav[3], dict):
                        nav_data = parsed_nav[3].get("data", parsed_nav[3])
                    elif isinstance(parsed_nav, dict):
                        nav_data = parsed_nav.get("data", parsed_nav)
                except Exception:
                    pass

            # 2. Featured Filters (e.g. pills, popular searches, toggles)
            if "BlockBuilderBrowseFeaturedFilters" in line:
                try:
                    parsed_feat = json.loads(payload_str)
                    def search_featured(obj: Any) -> Any:
                        if isinstance(obj, dict):
                            if "featuredFilters" in obj and isinstance(obj["featuredFilters"], list):
                                return obj["featuredFilters"]
                            if "filters" in obj and isinstance(obj["filters"], list):
                                if any(isinstance(x, dict) and x.get("__typename") == "FilterFeatured" for x in obj["filters"]):
                                    return obj["filters"]
                            for v in obj.values():
                                res = search_featured(v)
                                if res:
                                    return res
                        elif isinstance(obj, list):
                            for el in obj:
                                res = search_featured(el)
                                if res:
                                    return res
                        return None

                    ff = search_featured(parsed_feat)
                    if ff:
                        featured_data.extend(ff)
                except Exception:
                    pass

            # 3. Main Filter Refinements & Category Filter Blocks
            if any(term in line for term in [
                "BlockBuilderBrowseFilterRefinements",
                "BlockBuilderBrowseFilters",
                "BlockBuilderCategoryStandardExperience",
                '"refinements":',
                '"filters":['
            ]):
                try:
                    data = json.loads(payload_str)

                    def search_filters(obj: Any) -> Any:
                        if isinstance(obj, dict):
                            # Case A: GraphQL refinements.filters (e.g. BlockBuilderCategoryStandardExperience)
                            if "refinements" in obj and isinstance(obj["refinements"], dict) and "filters" in obj["refinements"]:
                                return obj["refinements"]["filters"]
                            # Case B: data.filters (e.g. BlockBuilderBrowseFilters)
                            if "filters" in obj and isinstance(obj["filters"], list):
                                if any(isinstance(x, dict) and ("filterId" in x or "options" in x or "id" in x or "subSections" in x or "buckets" in x) for x in obj["filters"]):
                                    return obj["filters"]
                            # Case C: legacy appliedFilters + filters
                            if "appliedFilters" in obj and "filters" in obj and isinstance(obj["filters"], list):
                                return obj["filters"]
                            for v in obj.values():
                                res = search_filters(v)
                                if res:
                                    return res
                        elif isinstance(obj, list):
                            for el in obj:
                                res = search_filters(el)
                                if res:
                                    return res
                        return None

                    fg = search_filters(data)
                    if fg and isinstance(fg, list) and fg not in raw_filter_blocks:
                        raw_filter_blocks.append(fg)
                except Exception:
                    pass

        return nav_data, featured_data, raw_filter_blocks

    def build_structured_data(
        self,
        category_name: str,
        category_url: str,
        nav_data: Optional[Dict[str, Any]],
        featured_data: Optional[List[Dict[str, Any]]],
        raw_filter_blocks: List[List[Dict[str, Any]]],
    ) -> List[Dict[str, str]]:
        """Transforms parsed flight blocks into normalized rows."""
        rows: List[Dict[str, str]] = []
        seen_filters: Set[Tuple[str, str, str]] = set()

        # 1. Category Navigation (if nav_data provided)
        if nav_data and isinstance(nav_data, dict) and self.include_nav:
            curr_cat = nav_data.get("displayName", category_name)
            curr_id = str(nav_data.get("categoryId", ""))
            curr_url = nav_data.get("url", category_url)
            breadcrumbs = [
                b.get("text", b.get("displayName", ""))
                for b in nav_data.get("displayBreadcrumbs", [])
            ]
            bc_str = " > ".join(breadcrumbs)

            parent_name = category_name
            parent_url = category_url
            raw_bcs = nav_data.get("displayBreadcrumbs", [])
            if len(raw_bcs) >= 2:
                parent_bc = raw_bcs[-2]
                parent_name = parent_bc.get("text", category_name)
                parent_url = parent_bc.get("url", category_url)

            rows.append({
                "category_name": parent_name,
                "category_url": parent_url,
                "hierarchy_path": f"Category Navigation > Current Category > {curr_cat}",
                "section": "Category Navigation",
                "filter_group": "Category",
                "sub_group": "Current Category",
                "attribute_name": curr_cat,
                "attribute_id": curr_id,
                "filter_id": "category",
                "input_type": "Category Current",
                "description": f"Breadcrumbs: {bc_str}",
                "tooltip": "",
                "range_min": "",
                "range_max": "",
                "range_unit": "",
                "is_enabled": "True",
                "is_selected": "True",
                "link_url": curr_url,
            })

            for sib in nav_data.get("siblings", []):
                sib_name = sib.get("displayName", "")
                sib_url = sib.get("url", "")
                if is_sale_category(sib_name, sib_url, self.ignore_sales):
                    continue
                rows.append({
                    "category_name": parent_name,
                    "category_url": parent_url,
                    "hierarchy_path": f"Category Navigation > Siblings > {sib_name}",
                    "section": "Category Navigation",
                    "filter_group": "Category",
                    "sub_group": "Sibling Categories",
                    "attribute_name": sib_name,
                    "attribute_id": str(sib.get("categoryId", "")),
                    "filter_id": "category",
                    "input_type": "Category Link",
                    "description": "",
                    "tooltip": "",
                    "range_min": "",
                    "range_max": "",
                    "range_unit": "",
                    "is_enabled": "True",
                    "is_selected": "False",
                    "link_url": sib_url,
                })

        # 2. Featured Filters
        if featured_data:
            for item in featured_data:
                if not isinstance(item, dict):
                    continue
                f_name = item.get("displayName") or item.get("name") or ""
                f_id = str(item.get("filterId") or item.get("id") or "")
                f_url = item.get("url", "")
                f_type = item.get("filterType") or item.get("type") or "Featured Pill"

                sig = ("Featured", "", f_name)
                if f_name and sig not in seen_filters:
                    seen_filters.add(sig)
                    rows.append({
                        "category_name": category_name,
                        "category_url": category_url,
                        "hierarchy_path": f"Featured Filters > {f_name}",
                        "section": "Featured Filters",
                        "filter_group": "Featured",
                        "sub_group": "Popular Searches",
                        "attribute_name": f_name,
                        "attribute_id": f_id,
                        "filter_id": f_id,
                        "input_type": f_type,
                        "description": "",
                        "tooltip": "",
                        "range_min": "",
                        "range_max": "",
                        "range_unit": "",
                        "is_enabled": "True",
                        "is_selected": "False",
                        "link_url": f_url,
                    })

        # 3. Main Filter Refinements (all blocks)
        for block in raw_filter_blocks:
            for f in block:
                if not isinstance(f, dict):
                    continue

                group_name = f.get("displayName") or f.get("name") or ""
                group_id = str(f.get("filterId") or f.get("id") or "")
                group_type = f.get("filterType") or f.get("type") or f.get("__typename") or ""

                # Format A: FilterSubsections (e.g. Color / Finish with subsections Color/Finish and Panel Ready)
                if "subSections" in f and isinstance(f["subSections"], list):
                    for sub in f["subSections"]:
                        if not isinstance(sub, dict):
                            continue
                        sub_name = sub.get("displayName") or sub.get("name") or ""
                        sub_id = str(sub.get("filterId") or sub.get("id") or group_id)
                        for opt in sub.get("options", []):
                            if not isinstance(opt, dict):
                                continue
                            opt_name = opt.get("displayName") or opt.get("name") or ""
                            opt_id = str(opt.get("filterOptionId") or opt.get("id") or "")
                            opt_selected = str(opt.get("isSelected") or opt.get("selected", False))
                            disp_info = opt.get("filterOptionDisplayInformation", {})
                            opt_enabled = str(disp_info.get("isEnabled", True) if isinstance(disp_info, dict) else not opt.get("disabled", False))

                            sig = (group_name, sub_name, opt_name)
                            if opt_name and sig not in seen_filters:
                                seen_filters.add(sig)
                                rows.append({
                                    "category_name": category_name,
                                    "category_url": category_url,
                                    "hierarchy_path": f"Filters > {group_name} > {sub_name} > {opt_name}",
                                    "section": "Filters",
                                    "filter_group": group_name,
                                    "sub_group": sub_name,
                                    "attribute_name": opt_name,
                                    "attribute_id": opt_id,
                                    "filter_id": sub_id,
                                    "input_type": "Checkbox",
                                    "description": "",
                                    "tooltip": "",
                                    "range_min": "",
                                    "range_max": "",
                                    "range_unit": "",
                                    "is_enabled": opt_enabled,
                                    "is_selected": opt_selected,
                                    "link_url": "",
                                })
                    continue

                # Format B: Single Toggle (FilterFeatured, e.g. Sale, Fast Delivery)
                if "option" in f and isinstance(f["option"], dict):
                    opt = f["option"]
                    opt_name = opt.get("displayName") or group_name
                    opt_id = str(opt.get("filterOptionId") or group_id)
                    opt_selected = str(opt.get("isSelected") or opt.get("selected", False))
                    disp_info = opt.get("filterOptionDisplayInformation", {})
                    opt_enabled = str(disp_info.get("isEnabled", True) if isinstance(disp_info, dict) else True)

                    sig = (group_name, "", opt_name)
                    if opt_name and sig not in seen_filters:
                        seen_filters.add(sig)
                        rows.append({
                            "category_name": category_name,
                            "category_url": category_url,
                            "hierarchy_path": f"Filters > {group_name}",
                            "section": "Filters",
                            "filter_group": group_name,
                            "sub_group": "",
                            "attribute_name": opt_name,
                            "attribute_id": opt_id,
                            "filter_id": group_id,
                            "input_type": "Toggle",
                            "description": "",
                            "tooltip": "",
                            "range_min": "",
                            "range_max": "",
                            "range_unit": "",
                            "is_enabled": opt_enabled,
                            "is_selected": opt_selected,
                            "link_url": "",
                        })

                # Format C: Standard Selectable Options (FilterWithSelectableOptions or legacy options list)
                if "options" in f and isinstance(f["options"], list):
                    for opt in f["options"]:
                        if not isinstance(opt, dict):
                            continue
                        opt_name = opt.get("displayName") or opt.get("name") or ""
                        opt_id = str(opt.get("filterOptionId") or opt.get("id") or "")
                        opt_selected = str(opt.get("isSelected") or opt.get("selected", False))
                        disp_info = opt.get("filterOptionDisplayInformation", {})
                        opt_enabled = str(disp_info.get("isEnabled", True) if isinstance(disp_info, dict) else not opt.get("disabled", False))
                        opt_url = opt.get("url", "")
                        opt_count = opt.get("count", "")
                        opt_desc = f"Count: {opt_count}" if opt_count != "" else ""

                        sig = (group_name, "", opt_name)
                        if opt_name and sig not in seen_filters:
                            seen_filters.add(sig)
                            rows.append({
                                "category_name": category_name,
                                "category_url": category_url,
                                "hierarchy_path": f"Filters > {group_name} > {opt_name}",
                                "section": "Filters",
                                "filter_group": group_name,
                                "sub_group": "",
                                "attribute_name": opt_name,
                                "attribute_id": opt_id,
                                "filter_id": group_id,
                                "input_type": "Checkbox",
                                "description": opt_desc,
                                "tooltip": "",
                                "range_min": "",
                                "range_max": "",
                                "range_unit": "",
                                "is_enabled": opt_enabled,
                                "is_selected": opt_selected,
                                "link_url": opt_url,
                            })

                # Format D: Numeric Ranges (FilterWithNumericRanges: buckets and range slider)
                if "buckets" in f and isinstance(f["buckets"], list):
                    for b in f["buckets"]:
                        if not isinstance(b, dict):
                            continue
                        b_name = b.get("displayName") or b.get("name") or ""
                        b_id = str(b.get("filterOptionId") or b.get("id") or "")
                        b_selected = str(b.get("isSelected") or b.get("selected", False))
                        disp_info = b.get("filterOptionDisplayInformation", {})
                        b_enabled = str(disp_info.get("isEnabled", True) if isinstance(disp_info, dict) else True)
                        r_start = str(b.get("optionRangeStart", "") if b.get("optionRangeStart") is not None else "")
                        r_end = str(b.get("optionRangeEnd", "") if b.get("optionRangeEnd") is not None else "")

                        sig = (group_name, "Buckets", b_name)
                        if b_name and sig not in seen_filters:
                            seen_filters.add(sig)
                            rows.append({
                                "category_name": category_name,
                                "category_url": category_url,
                                "hierarchy_path": f"Filters > {group_name} > {b_name}",
                                "section": "Filters",
                                "filter_group": group_name,
                                "sub_group": "Buckets",
                                "attribute_name": b_name,
                                "attribute_id": b_id,
                                "filter_id": group_id,
                                "input_type": "Bucket",
                                "description": f"{r_start} to {r_end}".strip(),
                                "tooltip": "",
                                "range_min": r_start,
                                "range_max": r_end,
                                "range_unit": f.get("units", ""),
                                "is_enabled": b_enabled,
                                "is_selected": b_selected,
                                "link_url": "",
                            })

                if ("range" in f and isinstance(f["range"], dict)) or group_type == "RANGE":
                    r = f.get("range", {})
                    r_min = str(r.get("minValue") if r.get("minValue") is not None else r.get("min", ""))
                    r_max = str(r.get("maxValue") if r.get("maxValue") is not None else r.get("max", ""))
                    r_unit = str(f.get("units") or r.get("unit", ""))

                    sig = (group_name, "Range", "Slider")
                    if sig not in seen_filters:
                        seen_filters.add(sig)
                        rows.append({
                            "category_name": category_name,
                            "category_url": category_url,
                            "hierarchy_path": f"Filters > {group_name} > Slider Range",
                            "section": "Filters",
                            "filter_group": group_name,
                            "sub_group": "Range",
                            "attribute_name": f"{group_name} Range",
                            "attribute_id": group_id,
                            "filter_id": group_id,
                            "input_type": "Range Slider",
                            "description": f"Min: {r_min}, Max: {r_max} {r_unit}".strip(),
                            "tooltip": "",
                            "range_min": r_min,
                            "range_max": r_max,
                            "range_unit": r_unit,
                            "is_enabled": "True",
                            "is_selected": "False",
                            "link_url": "",
                        })

        return rows

    def build_child_nav_row(
        self,
        parent_name: str,
        parent_url: str,
        child_name: str,
        child_url: str,
        hierarchy_path: Optional[str] = None,
    ) -> Dict[str, str]:
        """Creates a Category Navigation row for a child category under its parent."""
        cid_match = re.search(r"-c(\d+)\.html", child_url)
        cid = cid_match.group(1) if cid_match else ""
        h_path = hierarchy_path or f"Category Navigation > {parent_name} > {child_name}"

        return {
            "category_name": parent_name,
            "category_url": parent_url,
            "hierarchy_path": h_path,
            "section": "Category Navigation",
            "filter_group": "Category",
            "sub_group": "Subcategories",
            "attribute_name": child_name,
            "attribute_id": cid,
            "filter_id": "category",
            "input_type": "Category Link",
            "description": f"Subcategory of {parent_name}",
            "tooltip": "",
            "range_min": "",
            "range_max": "",
            "range_unit": "",
            "is_enabled": "True",
            "is_selected": "False",
            "link_url": child_url,
        }

    def process_category(
        self,
        name: str,
        url: str,
        filters_scraped_urls: Set[str],
        visited_landing_urls: Set[str],
        parent_chain: Optional[List[str]] = None,
        max_subcategories: Optional[int] = None,
    ) -> List[Dict[str, str]]:
        """
        Recursively processes categories down to the deepest leaf listings:
        1. Emits Category Navigation child rows under this parent.
        2. Extracts filters from this category if present.
        3. Recurses into child subcategories down to the deepest leaf node.
        """
        clean_url = url.split("?")[0]
        chain = (parent_chain or []) + [name]
        full_path_str = " > ".join(chain)
        rows: List[Dict[str, str]] = []

        print(f"\n========================================================")
        print(f"[*] Processing Category: {name}")
        print(f"    URL: {clean_url}")
        print(f"    Hierarchy: {full_path_str}")
        print(f"========================================================")

        try:
            html = self.fetch_page_html(clean_url)
        except Exception as e:
            print(f"[-] Failed to fetch {clean_url}: {e}")
            return []

        # 1. Discover Child Subcategories
        child_categories: List[Dict[str, str]] = []
        if "/sb0/" in clean_url:
            print(f"[+] Reached search browse listing (Leaf): {name} ({clean_url})")
        elif clean_url not in visited_landing_urls:
            visited_landing_urls.add(clean_url)
            child_categories = self.discover_subcategories(clean_url, html)

            if max_subcategories and len(parent_chain or []) == 0:
                child_categories = child_categories[:max_subcategories]
                print(f"[*] Limiting top-level execution to first {max_subcategories} subcategories.")

            if child_categories:
                print(f"[+] Discovered {len(child_categories)} subcategories inside '{name}':")
                for idx, c in enumerate(child_categories, 1):
                    print(f"    {idx}. {c['name']} -> {c['url']}")

                # Emit 1 Category Navigation row for EACH child under this parent
                for child in child_categories:
                    child_name = child["name"]
                    child_url = child["url"].split("?")[0]
                    nav_row = self.build_child_nav_row(
                        parent_name=name,
                        parent_url=clean_url,
                        child_name=child_name,
                        child_url=child_url,
                        hierarchy_path=f"Category Navigation > {full_path_str} > {child_name}",
                    )
                    rows.append(nav_row)
            else:
                print(f"[*] Leaf category reached (no child categories): {name}")

        # 2. Extract Direct Product Filters for THIS category page
        stream = self.extract_flight_stream(html)
        nav_data, featured_data, raw_filter_blocks = self.parse_flight_blocks(stream)

        if raw_filter_blocks and clean_url not in filters_scraped_urls:
            cat_filter_rows = self.build_structured_data(
                category_name=name,
                category_url=clean_url,
                nav_data=None,  # Navigation rows emitted cleanly above
                featured_data=featured_data,
                raw_filter_blocks=raw_filter_blocks,
            )
            print(f"[+] Extracted {len(cat_filter_rows)} filter rows directly for: {name}")
            rows.extend(cat_filter_rows)
            filters_scraped_urls.add(clean_url)
        elif clean_url in filters_scraped_urls:
            print(f"[*] Filters already collected for {name} ({clean_url}).")
        else:
            print(f"[*] No direct product filter blocks found on this page for: {name}")

        # 3. Recurse into children down to the deepest leaf node
        if child_categories:
            for child in child_categories:
                child_name = child["name"]
                child_url = child["url"].split("?")[0]

                # Avoid duplicate scraping if already visited
                if child_url in filters_scraped_urls and child_url in visited_landing_urls:
                    print(f"[*] {child_name} ({child_url}) already fully crawled; skipping re-scrape.")
                    continue

                if self.delay > 0:
                    time.sleep(self.delay)

                child_rows = self.process_category(
                    name=child_name,
                    url=child_url,
                    filters_scraped_urls=filters_scraped_urls,
                    visited_landing_urls=visited_landing_urls,
                    parent_chain=chain,
                )
                rows.extend(child_rows)

        return rows

    def export_excel_highlighted(self, rows: List[Dict[str, str]], filepath: str) -> None:
        """
        Exports rows to an Excel (.xlsx) file with background color highlights via openpyxl:
        - Human-readable title headers with spaces (Category Name, Category Url, etc.)
        - Category Navigation: soft blue
        - Featured Filters: soft pastel amber
        - Filter Groups: distinct alternating soft colors
        - Material: General, Leather, Fabric, Velvet, Wood, Metal
        - Color: Upholstery Color, Frame Color, Wood Tone
        - Category Boundary Separator: Solid black bottom border at the final row of each category!
        """
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("[!] openpyxl is not installed. Run: pip install openpyxl")
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Wayfair Filters"

        # Header styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        thin_side = Side(border_style="thin", color="D1D5DB")
        black_bottom_side = Side(border_style="medium", color="000000")

        default_cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        category_end_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=black_bottom_side)

        # Write Header with proper formatted names
        header_labels = [EXCEL_HEADER_MAPPING.get(fn, fn.replace("_", " ").title()) for fn in FIELDNAMES]
        ws.append(header_labels)
        for col_idx in range(1, len(FIELDNAMES) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = default_cell_border
        ws.row_dimensions[1].height = 28

        # Cache for grouping other filter categories
        group_color_cache: Dict[str, int] = {}

        # Write Data Rows with Highlighting and Category Separation Borders
        num_rows = len(rows)
        for i, row_data in enumerate(rows):
            row_idx = i + 2  # row 1 is header
            bg_hex, text_hex = get_row_color_scheme(row_data, group_color_cache)
            row_fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
            row_font = Font(name="Segoe UI", size=10, color=text_hex)

            row_values = [row_data.get(fn, "") for fn in FIELDNAMES]
            ws.append(row_values)

            # Check if this row is the final row of its category / section block
            curr_cat = row_data.get("category_name", "")
            curr_sec = row_data.get("section", "")
            next_row = rows[i + 1] if i + 1 < num_rows else None
            next_cat = next_row.get("category_name", "") if next_row else None
            next_sec = next_row.get("section", "") if next_row else None

            if curr_sec == "Category Navigation":
                is_cat_end = (next_sec != "Category Navigation")
            else:
                is_cat_end = (next_cat != curr_cat)

            row_border = category_end_border if is_cat_end else default_cell_border

            ws.row_dimensions[row_idx].height = 20
            for col_idx in range(1, len(FIELDNAMES) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = row_fill
                cell.font = row_font
                cell.border = row_border
                align_h = "left"
                if FIELDNAMES[col_idx - 1] in ["is_enabled", "is_selected", "attribute_id", "filter_id", "range_min", "range_max", "range_unit"]:
                    align_h = "center"
                cell.alignment = Alignment(
                    horizontal=align_h,
                    vertical="center",
                    wrap_text=(FIELDNAMES[col_idx - 1] in ["description", "tooltip"]),
                )

        # Freeze Header Row
        ws.freeze_panes = "A2"

        # Auto-adjust column widths
        for col_idx, col_name in enumerate(FIELDNAMES, start=1):
            col_letter = get_column_letter(col_idx)
            header_label = EXCEL_HEADER_MAPPING.get(col_name, col_name.replace("_", " ").title())
            max_len = len(header_label)
            for row_data in rows[:150]:
                val = str(row_data.get(col_name, "") or "")
                max_len = max(max_len, min(len(val), 40))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

        wb.save(filepath)
        print(f"[+] Successfully exported highlighted Excel file: {filepath}")

    def export_csv(self, rows: List[Dict[str, str]], filepath: str) -> None:
        """Exports clean, tabular rows to a standard CSV file."""
        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
            writer.writeheader()
            writer.writerows(rows)
        print(f"[+] Successfully exported {len(rows)} rows to CSV: {filepath}")

    def run(
        self,
        urls: Optional[List[str]] = None,
        excel_path: str = DEFAULT_EXCEL_OUTPUT,
        csv_path: Optional[str] = DEFAULT_CSV_OUTPUT,
        max_categories: Optional[int] = None,
    ) -> List[Dict[str, str]]:
        """
        Main runner: Iterates through all provided seed URLs, crawling each
        down to the leaf nodes and accumulating all rows into Excel and CSV.
        """
        target_urls = urls or DEFAULT_SEED_URLS
        if not target_urls:
            print("[!] No URLs provided to scrape. Specify URLs in DEFAULT_SEED_URLS or pass via --urls / --urls-file.")
            return []

        filters_scraped_urls: Set[str] = set()
        visited_landing_urls: Set[str] = set()
        all_rows: List[Dict[str, str]] = []

        print(f"\n========================================================")
        print(f"[*] Starting Multi-URL Scraping Job: {len(target_urls)} seed URL(s)")
        print(f"========================================================")

        for idx, url in enumerate(target_urls, 1):
            clean_url = url.split("?")[0].strip()
            if not clean_url:
                continue

            cat_name = clean_url.split("/")[-1].split("-c")[0].replace("-", " ").title()
            print(f"\n>>> [Seed {idx}/{len(target_urls)}] Processing: {cat_name} -> {clean_url}")

            url_rows = self.process_category(
                name=cat_name,
                url=clean_url,
                filters_scraped_urls=filters_scraped_urls,
                visited_landing_urls=visited_landing_urls,
                parent_chain=[],
                max_subcategories=max_categories,
            )
            all_rows.extend(url_rows)

        print(f"\n========================================================")
        print(f"[+] Total accumulated rows across all seed URLs: {len(all_rows)}")
        print(f"========================================================")

        # 1. Export highlighted Excel (.xlsx)
        if excel_path and all_rows:
            self.export_excel_highlighted(all_rows, excel_path)

        # 2. Export clean CSV
        if csv_path and all_rows:
            self.export_csv(all_rows, csv_path)

        return all_rows


# ==============================================================================
# CLI ARGUMENT PARSER
# ==============================================================================
def parse_args():
    parser = argparse.ArgumentParser(
        description="Scrape Wayfair filter and category hierarchy data across multiple seed URLs down to leaf nodes."
    )
    parser.add_argument(
        "--urls",
        nargs="+",
        default=None,
        help="List of seed URLs separated by spaces. Example: --urls https://www.wayfair.com/... https://www.wayfair.com/...",
    )
    parser.add_argument(
        "--urls-file",
        default=None,
        help="Path to a text file containing one seed URL per line.",
    )
    parser.add_argument(
        "--output-excel",
        default=DEFAULT_EXCEL_OUTPUT,
        help=f"Output Excel (.xlsx) path (default: {DEFAULT_EXCEL_OUTPUT})",
    )
    parser.add_argument(
        "--output-csv",
        default=DEFAULT_CSV_OUTPUT,
        help=f"Output CSV path (default: {DEFAULT_CSV_OUTPUT})",
    )
    parser.add_argument(
        "--max-categories",
        type=int,
        default=None,
        help="Limit number of child categories to scrape per seed URL (useful for testing).",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=1.0,
        help="Delay in seconds between category page requests (default: 1.0s).",
    )
    parser.add_argument(
        "--zyte-key",
        default=os.getenv("ZYTE_API_KEY") or DEFAULT_ZYTE_KEY,
        help="Zyte API key for anti-blocking proxy.",
    )
    parser.add_argument(
        "--no-zyte",
        action="store_true",
        help="Disable Zyte API and make direct HTTP requests.",
    )
    parser.add_argument(
        "--ignore-sales",
        action="store_true",
        help="Filter out sale/clearance/deals categories.",
    )
    return parser.parse_args()


def load_urls_from_file(filepath: str) -> List[str]:
    """Reads URLs from a plain text file, ignoring empty lines and comments (#)."""
    urls = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            stripped = line.strip()
            if stripped and not stripped.startswith("#"):
                urls.append(stripped)
    return urls


def main():
    args = parse_args()

    # Determine seed URLs: CLI --urls > CLI --urls-file > DEFAULT_SEED_URLS in code
    seed_urls: List[str] = []
    if args.urls:
        seed_urls = args.urls
    elif args.urls_file:
        if not os.path.exists(args.urls_file):
            print(f"[!] Error: URLs file '{args.urls_file}' not found.")
            sys.exit(1)
        seed_urls = load_urls_from_file(args.urls_file)
    else:
        seed_urls = DEFAULT_SEED_URLS

    if not seed_urls:
        print("[!] No URLs provided. Either add them to DEFAULT_SEED_URLS in the script or pass via --urls / --urls-file.")
        sys.exit(1)

    scraper = WayfairMultiFilterScraper(
        zyte_key=args.zyte_key,
        use_zyte=not args.no_zyte,
        include_nav=True,
        ignore_sales=args.ignore_sales,
        delay=args.delay,
    )

    scraper.run(
        urls=seed_urls,
        excel_path=args.output_excel,
        csv_path=args.output_csv,
        max_categories=args.max_categories,
    )


if __name__ == "__main__":
    main()
